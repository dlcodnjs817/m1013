# -*- coding: utf-8 -*-
"""2026 한국생산제조학회 추계학술대회 논문 — 기존 연구 대비 차별점 정리 (2026-09-11)"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x33, 0x64)
ACC  = RGBColor(0xC0, 0x50, 0x2A)
GRAY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---- 기본 스타일 / 여백 ----
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.35

for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)

def _run(p, text, size=10.5, bold=False, color=None, italic=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r

def shade(el, hexcolor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    el.get_or_add_tcPr().append(sh) if el.tag.endswith("tc") else el.get_or_add_pPr().append(sh)

def bar(text, sub=None):
    """남색 바 제목"""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c._tc, "1F3364")
    c.width = Cm(16.2)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    if sub:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        _run(p2, sub, size=9.5, color=RGBColor(0xC8, 0xD4, 0xE8))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=11.5, bold=True, color=NAVY)

def body(text, indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    _run(p, text)
    return p

def bullet(text, indent=0.5, mark="•"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent + 0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(3)
    _run(p, f"{mark} {text}")
    return p

def rich_bullet(segs, indent=0.5, mark="•"):
    """segs = [(text, bold, color), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent + 0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(3)
    _run(p, f"{mark} ")
    for text, b, col in segs:
        _run(p, text, bold=b, color=col)
    return p

def flowbox(title, text, fill="EEF2F8", border="1F3364"):
    """강조 박스"""
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c._tc, fill)
    tcPr = c._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "12" if edge == "left" else "4")
        e.set(qn("w:color"), border)
        borders.append(e)
    tcPr.append(borders)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    if title:
        _run(p, title, size=10.5, bold=True, color=NAVY)
        p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    _run(p, text, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c._tc, "1F3364")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        _run(p, htxt, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _run(p, v, size=9.5, bold=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run(p, v, size=9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)



# =====================================================================
# 표지
# =====================================================================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
_run(p, "학회 논문 차별점 정리", size=18, bold=True, color=NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
_run(p, "저가 5축 매니퓰레이터 시연 데이터의 6축 협동로봇 모방학습 전이", size=11, color=GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
_run(p, "2026 한국생산제조학회 추계학술대회 (12. 9 ~ 12. 12, 쉐라톤 제주)", size=10, color=GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_after = Pt(12)
_run(p, "2026. 09. 11.", size=10, color=GRAY)

# =====================================================================
# 0. 한 문단 요약
# =====================================================================
bar("0. 지도교수님 보고용 요약")
flowbox("■ 한 문단",
    "기존 이종형상(cross-embodiment) 연구는 로봇 간 차이를 '행동 차원 불일치'로 보고 최대 자유도 패딩·마스킹·"
    "embodiment 토큰으로 처리합니다. 본 연구의 소스(OMX 5축)와 목표(M1013 6축)는 EE pose 로 올리면 차원이 "
    "이미 같고(7D↔7D), 남는 차이는 차원이 아니라 '소스가 도달할 수 있는 부분집합의 모양'입니다. 오늘 78,009 "
    "프레임 전수 계산으로 이 부분집합이 닫힌형 구속식 1개(EE x축 방위각 = EE 위치 방위각, 잔차 0.21°)로 정확히 "
    "적히는 SE(3)의 5차원 부분다양체임을 확인했습니다. 축이 '없는' 것이 아니라 '있는데 종속'인 경우라 마스킹으로는 "
    "표현조차 되지 않으며, 이 결손이 스칼라로 주어지므로 전이 성능을 결손 크기의 함수로 측정할 수 있습니다. "
    "여기에 EE pose 공통 행동공간의 대가(IK 특이점·작업영역 제한)를 정량 보고하는 것, 그리고 저가 데모의 값어치를 "
    "'협동로봇 실물 데모 N개 등가'로 환산하는 것이 기존 연구와 갈리는 지점입니다.")

h2("제목안")
table(["구분", "제목"], [
    ["추천 (학회 청중 맞춤)", "저가 5축 매니퓰레이터 시연 데이터의 6축 협동로봇 모방학습 전이: 행동 다양체 구속 관점의 정량 평가"],
    ["영문", "Transferring Low-Cost 5-DOF Manipulator Demonstrations to an 6-DOF Collaborative Robot: A Quantitative Study under Action-Manifold Constraints"],
    ["대안 (학술 강조)", "저자유도 데모는 언제 값어치를 하는가: 행동 다양체 구속 관점의 이종형상 VLA 전이"],
], widths=[3.4, 12.8])
body("한국생산제조학회는 로봇러닝 전문 학회가 아니므로 '저가 로봇 데모를 협동로봇에 재활용'이라는 실용 축을 "
     "제목 앞에 두고, '행동 다양체 구속'은 방법론의 핵심어로 뒤에 붙이는 것을 권합니다.")

doc.add_page_break()

# =====================================================================
# 1. 지금까지 한 것
# =====================================================================
bar("1. 지금까지 한 것 — 논문에 쓸 수 있는 자산")

h2("1.1  데이터 · 표현 · 재타겟팅")
table(["항목", "내용", "수치"], [
    ["소스 데모", "OMX(omx_f, 5축+그리퍼) 실물 pick&place, 손목캠+정면캠 640×480 30fps", "163 ep / 78,009 fr (v6: 159 ep / 76,345 fr)"],
    ["공통 표현", "URDF FK 로 joint → EE pose 7D [x,y,z,rx,ry,rz,grip]", "OMX·M1013 모두 7D"],
    ["재타겟팅 IK", "자체 DLS IK: 오차비례 감쇠 + 시드 체인 + 분지 거부 + 손목 45° 보정", "159/159 성공, 실패·분지 0"],
    ["재타겟팅 품질", "프레임당 최대 Δq / FK 역검증 위치·각도 오차", "2.95° / ≤4.89 mm / ≤1.03°"],
    ["특이점 여유", "손목 특이점(q5=0)까지의 최소 거리", "min|q5| = 45.97°"],
    ["작업대 배치", "54후보 → 5후보 → 확정 스윕(min|q5|, j3 여유, maxΔq 채점)", "OFFSET (+0.05, −0.15, +0.10)"],
], widths=[2.8, 8.4, 5.0])

h2("1.2  관측공간 정합 (카메라)")
table(["항목", "내용", "수치"], [
    ["Isaac 씬 정합", "OMX v6 학습영상과 같은 구도의 카메라 2대를 M1013 씬에 장착", "cameras_v6match.json"],
    ["손목캠 해석 배치", "제작 가능 영역 전수 탐색 + 렌즈 스탠드오프 ±12 mm 오차 포함 여유로 결정 (09-11 확정)", "플랜지 로컬 광심 (0,−65,−10) mm, 틸트 21.8°, 세로 장착, 여유 0.63"],
    ["화각 요구조건", "파지 정렬구간(2.0~0.7 s 전) 내내 큐브 8꼭짓점 가시 여부로 판정 (※ 구 마운트 기준 — 최종 마운트로 재실행 필요)", "63.4° 불합격(2/8 ep 이탈) / 85.6° 합격(8/8)"],
    ["검증기", "Isaac 없이 FK 만으로 투영하는 순수기하 검증기", "wristcam_check.py, 수 초"],
], widths=[2.8, 8.4, 5.0])

h2("1.3  실물 시스템 (M1013 + MHF2-16D2 이관)")
table(["항목", "내용"], [
    ["연결부 검증", "어댑터·핑거·브래킷을 cadquery 커널로 전수 검증 — 조인트 5종 동축 일치, 간섭 0, ISO 9409-1-50-4-M6 플랜지 도면 대조 완료"],
    ["기하 정본", "TCP = 파지점 0.0725 m, Isaac 상판 0.4624 m, npz 가 실물 그리퍼 기하를 그대로 적재"],
    ["실기 SW 4종", "camera_uvc_node(1280×720→960×720 크롭→640×480) / gripper_ctl(히스테리시스+grip-lead) / measure_grip_lead / replay_v6_real(v7 수집)"],
], widths=[3.0, 13.2])

h2("1.4  오늘(09-11) 신규 확인 — 논문의 핵심 수치")
body("action_manifold_analysis.py (tx90 저장소) · 163 에피소드 / 78,009 프레임 전수 · omx_f.urdf FK")
table(["검증 대상", "잔차 RMS", "최대", "해석"], [
    ["az(tool x축) − joint1", "0.00000°", "0.00000°", "정확한 항등식 — EE 방위 성분의 원천은 q1 뿐"],
    ["az(EE 위치) − joint1", "0.20953°", "0.50766°", "횡오프셋 0.66 mm 환산 — 기구 오프셋"],
    ["az(tool x축) − az(EE 위치)", "0.20953°", "0.50766°", "SE(3) 위 닫힌형 홀로노믹 구속 1개"],
], widths=[4.4, 2.4, 2.4, 7.0])
table(["지표", "값", "의미"], [
    ["6D pose 유효랭크 (엔트로피)", "2.70 / 6", "실데이터가 차지하는 유효 차원"],
    ["6D 주성분 설명비율", "0.636 / 0.196 / 0.143 / 0.021 / 0.004 / 0.001", "4번째 축부터 급감"],
    ["q1(=yaw) 스팬", "16.5° (물리 리밋 ±180°)", "기구 한계보다 데이터가 훨씬 좁음"],
    ["tool z 고도각", "13.6° ~ 83.9° (std 10.6°)", "접근 틸트 커버리지"],
], widths=[4.6, 6.0, 5.6])
flowbox("■ 기존 연구방향 문서(09-09)와 달라진 점",
    "09-09 문서는 'yaw ≈ atan2(y, x)' 로 근사 서술했으나, 정확한 형태는 'EE x축의 방위각 = EE 위치의 방위각' 이며 "
    "이는 근사가 아니라 5축 직렬체인의 기구학적 항등식입니다. 7D 표현 중 실질 자유도는 6 (pose 5 + grip 1) 이고, "
    "도달집합은 SE(3)의 5차원 부분다양체입니다. 논문에는 이 정확한 형태로 기술해야 합니다.",
    fill="FBF0EC", border="C0502A")

doc.add_page_break()

# =====================================================================
# 2. 기존 연구 지형
# =====================================================================
bar("2. 기존 연구 지형 — 무엇이 이미 되어 있는가")
table(["계열", "대표 연구", "축 불일치 처리 방식", "본 연구와의 관계"], [
    ["대규모 이종 데이터 사전학습", "Open X-Embodiment / RT-X, Octo, OpenVLA, CrossFormer, HPT",
     "최대 DOF 패딩 + 마스킹 + embodiment 토큰. '차원' 문제로 취급", "소스·타깃 모두 full-DOF(Franka 7, UR5e 6, WidowX 6). DOF 결손 소스 없음"],
    ["형상 정보 주입", "Embedding Morphology into Transformers (2603.00182)",
     "kinematic 토큰 + 토폴로지 어텐션 + FiLM", "형상을 '입력'으로 넣지만 도달집합의 결손을 다루진 않음"],
    ["통제 실험", "ZETA (2609.02546)", "EEF-Delta 표현이 절대좌표 대비 +15%p. appearance/gripper/arm/full 4분류",
     "가장 가까운 통제 연구. 명시적으로 DOF 결손·데이터효율곡선 미수행"],
    ["데이터 효율", "Data Analogies (2603.06450)", "소스 다양성 ↑ → 타깃 성공률 곡선, 50-shot 타깃 대비",
     "곡선은 그리나 '타깃 데모 N개 등가' 환산은 없음"],
    ["시각 갭 완화", "RoVi-Aug, Mirage, Cloak (2606.22836)", "로봇 외형·시점 augmentation / EE 마스킹",
     "시각을 '가린다'. 본 연구는 반대로 타깃 리그를 소스 관측분포에 맞춰 '설계'"],
    ["저가 로봇 IL", "SmolVLA, SO-100/101, LeRobot 생태계", "저가 암 위에서 학습·평가",
     "저가 데모가 협동로봇에 값어치를 하는지는 미답"],
    ["행동공간 통일", "Unify Actions in Camera Frame, PointAction, Latent Action Diffusion",
     "카메라 프레임·3D 점·잠재 행동으로 통일", "통일 표현의 '대가'(IK 특이점·작업영역)는 보고하지 않음"],
], widths=[2.8, 4.2, 4.6, 4.6])

doc.add_page_break()

# =====================================================================
# 3. 차별점
# =====================================================================
bar("3. 기존 연구와 다른 점", "6가지 — 성립 조건별로 구분")

def diff(no, title, prior, ours, why, cond, strong=False):
    h2(f"차별점 {no}.  {title}")
    rich_bullet([("기존: ", True, None), (prior, False, None)])
    rich_bullet([("본 연구: ", True, ACC), (ours, False, None)])
    rich_bullet([("왜 새로운가: ", True, None), (why, False, None)])
    rich_bullet([("성립 조건: ", True, None), (cond, False, GRAY)])

diff(1, "'축 개수 불일치'가 아니라 '행동 다양체의 구속'으로 문제를 재정의",
    "OXE·Octo·OpenVLA·CrossFormer·HPT 는 이종 로봇을 최대 DOF 패딩+마스킹+embodiment 토큰으로 처리 — 차원(dimension) 문제로 봄.",
    "EE pose 로 올리면 차원은 이미 7D↔7D 로 같음. 남는 것은 소스의 도달집합이 SE(3) 안에서 어떤 모양의 부분집합이냐이며, 이 경우 닫힌형 구속식 1개로 정확히 적힘.",
    "패딩·마스킹은 '없는 축'을 다룸. 여기서는 축이 없는 게 아니라 있는데 종속임 — 마스킹으로는 표현조차 되지 않음.",
    "지금 데이터로 성립 (78,009 프레임 검증 완료).", strong=True)

diff(2, "결손이 스칼라로 측정 가능한 드문 세팅",
    "대부분 소스·타깃이 둘 다 full-DOF. ZETA 가 가장 통제된 연구인데 DOF 결손 소스는 다루지 않음. 형상 격차가 '다르다'까지는 가도 '무엇이 얼마나 부족한지'를 적지 못함.",
    "결손이 구속식 1개로 정확히 주어지므로, 전이 성능을 '평가 태스크가 요구하는 yaw 이탈각'의 함수로 그릴 수 있음. 통제 실험이 아니라 측정 실험이 됨.",
    "형상 격차를 연속 변수로 다룬 사례가 없음. yaw 불필요 태스크 vs yaw 필요 태스크의 성능 격차 = 자유도 결손의 비용.",
    "실물 결과 필요 (v7 수집·이어학습·평가군 2종).")

diff(3, "공통 행동공간의 '청구서'를 함께 보고",
    "EE-pose / EEF-delta 공통 행동공간은 사실상 표준(ZETA: 절대좌표 대비 +15%p). 그러나 정책 출력 pose 를 타깃 관절로 되돌릴 때의 IK 특이점·분지 점프·작업영역 제한을 정량 보고한 사례가 사실상 없음.",
    "재타겟팅 159/159, 프레임당 최대 Δq 2.95°, FK 역검증 ≤4.89 mm, min|q5| 45.97°, manipulability 기반 작업영역 확정 절차 — 비용을 숫자로 제시.",
    "'전이가 된다'가 아니라 '이 비용을 치르면 이만큼 된다'로 보고하는 것 자체가 생산제조 청중에게 실용적 기여.",
    "지금 데이터로 성립.", strong=True)

diff(4, "관측공간 정합을 사전조건으로 명시하고 기하로 푼다",
    "RoVi-Aug·Mirage·Cloak 은 로봇 외형·시점을 augmentation 이나 마스킹으로 덮음.",
    "반대로 타깃 리그(손목캠 위치·틸트·HFOV)를 정량 판정으로 설계 — '파지 정렬구간 내내 대상이 프레임 안' + 렌즈 스탠드오프 오차 포함 여유. 화각 요구조건 63.4° 불합격 / 85.6° 합격, 최종 마운트 광심 (0,−65,−10)·틸트 21.8°.",
    "'이종형상 전이를 하려면 하드웨어를 어떻게 설계해야 하는가'에 대한 처방. 알고리즘 논문이 다루지 않는 층위이며 생산제조 학회 성격에 맞음.",
    "지금 데이터로 성립 (wristcam_check.py 결과).", strong=True)

diff(5, "형상 격차 vs 도메인 격차의 2×2 분리",
    "sim-to-real 격차를 perception / embodiment / controller / physics 로 나누는 서술은 있으나, '형상 다르고 실물'(OMX) 과 '형상 같고 시뮬'(Isaac M1013) 을 같은 타깃·같은 태스크에서 같은 저울에 올린 통제 실험은 찾지 못함.",
    "OMX 실물 159 ep 사전학습 vs Isaac M1013 자체 데모 사전학습 → 같은 M1013 실물 평가군에서 비교.",
    "'형상 격차와 도메인 격차 중 무엇이 VLA 전이에 더 해로운가'라는 독립 질문에 답할 수 있는 구조.",
    "실물 결과 + Isaac 데모 생성 필요. 후속 논문 후보.")

diff(6, "저가 5축 → 6축 협동로봇이라는 실용 축과 '등가 데모 수' 지표",
    "SO-100/101·OMX 등 저가 암 모방학습은 급증했으나, 그 데이터가 협동로봇에 값어치를 하는지는 답이 없음. 데이터 효율 곡선은 흔하나 '타깃 실물 데모 N개 등가'로 환산해 보고하는 사례는 찾기 어려움.",
    "x축 M1013 실물 데모 수(10/30/60/120), y축 성공률, 곡선 3종(scratch / 비전만 전이 / 전체 전이) → 'OMX 159 ep = M1013 실물 데모 N개' 한 문장.",
    "생산 현장의 질문('교육용 로봇으로 찍은 데이터를 협동로봇에 써도 되나')에 직접 답하는 지표.",
    "실물 결과 필요.")

flowbox("■ 성립 조건 요약",
    "차별점 1·3·4 는 지금 보유한 데이터만으로 이미 성립합니다 (분석·시스템 성격). "
    "차별점 2·6 은 v7 수집과 이어학습 결과가 있어야 성립하며, 이것이 논문의 '결과' 절이 됩니다. "
    "차별점 5 는 Isaac 데모 생성이 추가로 필요하므로 이번 학회에서는 향후 연구로 두는 것을 권합니다.",
    fill="EEF2F8")

doc.add_page_break()

# =====================================================================
# 4. 기여 문장
# =====================================================================
bar("4. 기여 문장(Contributions) 초안")
bullet("5축 매니퓰레이터 시연 데이터의 EE pose 도달집합이 SE(3)의 5차원 부분다양체이며 그 구속이 닫힌형 1개 식(EE x축 방위각 = EE 위치 방위각)으로 주어짐을 실데이터 78,009 프레임으로 검증하였다 (잔차 RMS 0.21°).", mark="①")
bullet("이 구속을 제거 대상이 아닌 통제 변수로 삼아, 소스 데모의 전이 기여도를 '목표 로봇 실물 데모 등가 수'로 정량화하는 데이터 효율 실험 설계를 제안하고 M1013 실물에서 측정하였다.", mark="②")
bullet("EE pose 공통 행동공간이 수반하는 IK 재타겟팅 비용(특이점 여유, 프레임 간 관절 변화, 역검증 오차, 작업영역 제한)을 정량 보고하여, 전이 이득을 비용과 함께 평가하는 틀을 제시하였다.", mark="③")
bullet("이종형상 전이의 사전조건으로 관측공간 정합을 두고, 손목 카메라의 위치·틸트·화각을 기하 판정으로 결정하는 절차와 그 결과(HFOV 85.6°)를 제시하였다.", mark="④")
body("② 는 결과가 나온 뒤 수치를 채워 넣습니다. 결과가 초록 마감에 못 미치면 ②를 '실험 설계 제안'으로 낮추고 ①③④ 로 분석·시스템 논문을 구성합니다.")

# =====================================================================
# 5. 일정
# =====================================================================
bar("5. 일정과 리스크")
table(["시점", "내용", "비고"], [
    ["09-14 (월)", "M1013 실물 도착 + 절삭 외주 2건 발주", "9/20 에서 앞당겨짐. 발주서·STEP·DXF 는 준비 완료"],
    ["09-14 ~ 금구 입고", "배선·ROS2 연결·카메라 노드·manipulability 실측·작업대 확정·저속 궤적 재생", "그리퍼 없이 가능한 작업"],
    ["금구 입고 (09-21 ~ 09-28 예상, 5~10 영업일)", "Φ6 H7 현장 리머 → 그리퍼 장착 → grip-lead 실측 → LEAD 확정", "★ 새 병목. 납기 회신 확인 필요"],
    ["금구 입고 + 1주", "v7 실기 수집 30~50 ep (yaw 다양성 포함, 평가군 2종 분리)", ""],
    ["+ 2주", "이어학습 3종 × 데모 수 4단계, 평가", "GPU 시간 산정 필요"],
    ["10월 말 ~ 11월 초 (확인 필요)", "초록 마감 (학회 관례. 홈페이지에 아직 미공지)", "★ 마감일 확인 최우선"],
    ["12-09 ~ 12-12", "학술대회 발표, 쉐라톤 제주", ""],
], widths=[4.0, 8.2, 4.0])

flowbox("■ 지금 해야 할 것",
    "① 학회 초록 마감일·원고 형식(초록 분량, 구두/포스터) 확인.  ② 절삭 외주 납기 회신 확인 — 이게 v7 수집 착수일을 결정.  "
    "③ 현재 v6 모델은 159 ep 전부를 학습에 써서 검증셋이 없음 → 논문용 수치는 split 재학습으로 다시 뽑아야 함.  "
    "④ 09-09 연구방향 문서의 'yaw ≈ atan2(y,x)' 서술을 정확한 구속식으로 정정.",
    fill="FBF0EC", border="C0502A")

doc.add_page_break()
# =====================================================================
# 6. 논문 구성안 (결과 없는 판)
# =====================================================================
bar("6. 논문 구성안 — 현재 진행 상태 기준 (결과 없는 판)",
    "실물 결과가 나오면 6절만 갈아끼운다. 2·3·4·5절은 어느 판이든 그대로")

h2("제목")
body("5축 시연 데이터의 6축 협동로봇 전이를 위한 행동 다양체 분석과 EE pose 재타겟팅의 비용 정량화")
body("'모방학습 전이'를 제목에서 빼고 '전이를 위한'으로 쓴 것은 의도적 — 리뷰어가 전이 성능 표를 찾지 않게 한다.")

h2("논리 흐름")
flowbox(None,
    "저가 5축 데모를 협동로봇에 쓰고 싶다 → 기존은 축 수 차이를 패딩·마스킹으로 메운다 → EE pose 로 올리면 차원은 같고, "
    "진짜 격차는 소스 도달집합이 SE(3)의 5차원 부분다양체라는 것이다(닫힌형 구속식) → 그 표현을 쓰면 IK 재타겟팅 비용이 "
    "생기는데 그 비용을 측정했다 → 관측 쪽도 기하로 맞췄다 → 시뮬 재생으로 파이프라인이 도는 것을 확인했다 → "
    "실물 데이터 효율 곡선은 향후 연구.", fill="EEF2F8")

h2("절별 내용")
table(["절", "내용", "근거 (수치·그림·표)"], [
    ["초록", "문제(축 수) → 재정의(부분다양체) → 비용 정량화 → 정합·검증", "—"],
    ["1. 서론", "저가 암 IL 보편화 vs 협동로봇 재수집 비용. 기존은 '없는 축'을 다루는 방법. 우리는 '있는데 종속'. 기여 ①~④", "OXE/Octo/OpenVLA/CrossFormer/HPT, ZETA"],
    ["2. 관련 연구", "이종형상 사전학습 / 통제 실험 / 시각 갭 / 저가 로봇 IL / 행동공간 통일 — 각 두 줄. 공통 결론: DOF 결손 소스·재타겟팅 비용 미답", "2절 표 참조"],
    ["3. 행동 다양체 분석", "3.1 세팅 · 3.2 공통 표현 a=[p,r,g]∈R⁷ · 3.3 구속식 유도 식(1) atan2(p_y,p_x)=atan2((Re₁)_y,(Re₁)_x)=q₁ → 여차원 1 부분다양체 · 3.4 실데이터 검증 · 3.5 커버리지(결손 2층위: 기구적 + 데이터 분포) · 3.6 함의",
     "표 1 잔차 0.00000°/0.21°/0.51°, 횡오프셋 0.66 mm · 그림 1 방위각 산점도 · 그림 2 특이값 스펙트럼 · 유효랭크 2.70/6 · q₁ 스팬 16.5°"],
    ["4. 재타겟팅 파이프라인과 비용", "4.1 7MA→6Hz 앵커→오프셋→TCP 0.0725→손목 45°→DLS IK(λ∝오차·시드·분지거부) · 4.2 배치 스윕 54→5→1 · 4.3 비용 지표 5종 정의(제안) · 4.4 결과 · 4.5 협동로봇 맥락 = 안전 요건",
     "표 2 159/159, Δq 2.95°, 4.89 mm/1.03°, min|q₅| 45.97°, 분지 0 · 그림 3 스윕 맵 · 그림 4 에피소드별 분포 · TCP 미보정 시 84 mm"],
    ["5. 관측공간 정합", "5.1 원칙(구도 먼저) · 5.2 손목캠 기하 결정(전수 탐색 + 스탠드오프 오차) · 5.3 화각 요구조건 · 5.4 정면캠 동일 기종 이관",
     "광심 (0,−65,−10)·틸트 21.8°·HFOV 85.6°·여유 0.63 · 63.4° 불합격/85.6° 합격 · 그림 5 접근 구간 가시성 (★최종 마운트로 재실행)"],
    ["6. 시뮬레이션 검증", "Isaac Sim 5.1, 실물 그리퍼 기하, 상판 0.4624. 재타겟팅 궤적 재생. (추가 권장) 검증셋 분리 ACT 재학습 → held-out 오차 + 정책 출력 재생",
     "표 3 ep0/50/130 성공, 배치오차 7.5/1.6/5.2 mm · (추가 시) held-out open-loop 오차"],
    ["7. 한계·향후", "제한 작업영역, 접촉 물리, 실물 미검증 / 실물 데이터 효율 곡선·yaw 평가군·등가 데모 수·형상 vs 도메인 2×2", "—"],
    ["8. 결론", "3문장", "—"],
], widths=[2.6, 8.2, 5.4])

h2("그림·표 목록 (전부 현재 데이터로 제작 가능)")
table(["번호", "내용", "출처"], [
    ["그림 1", "az(Re₁) vs az(p) 산점도, 78,009점이 y=x 위", "action_manifold_analysis.py"],
    ["그림 2", "6D pose 특이값 스펙트럼 (4번째 축부터 급감)", "action_manifold_analysis.py"],
    ["그림 3", "작업대 배치 스윕 맵 (54 후보 채점)", "sweep_result.json"],
    ["그림 4", "에피소드별 min|q₅| · maxΔq 분포", "v6_staging/report.json"],
    ["그림 5", "손목캠 접근 구간 큐브 꼭짓점 가시성, 63.4° vs 85.6°", "wristcam_check.py (최종 마운트로 재실행)"],
    ["표 1", "구속식 잔차 3종 + 횡오프셋 환산", "action_manifold_analysis.py"],
    ["표 2", "재타겟팅 비용 지표 5종", "report.json"],
    ["표 3", "Isaac 재생 성공·배치오차", "replay_isaac.py"],
], widths=[1.8, 9.4, 5.0])

flowbox("■ 주장 범위",
    "이 판에서는 '전이가 된다'를 주장하지 않는다. 주장할 수 있는 것은 '무엇이 진짜 격차이고, 그 격차를 넘는 데 비용이 얼마인가'까지다. "
    "그걸 정확히 주장하면 결과 없이도 논문이 선다. 6절에 검증셋 분리 ACT 재학습을 넣으면 '재타겟팅 데이터로 학습한 정책이 "
    "타깃 기구학에서 실행 가능'까지 한 단계 올라간다 — 로봇 없이 컨테이너에서 지금 가능.",
    fill="FBF0EC", border="C0502A")

out = "/home/kim/m1013/대화록 및 PPT/M1013_학회논문_차별점_20260911.docx"
doc.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
