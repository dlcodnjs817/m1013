# -*- coding: utf-8 -*-
"""M1013 VLA 전이학습 연구방향 정리 워드 문서 생성 (2026-09-09)"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x33, 0x64)
ACC  = RGBColor(0xC0, 0x50, 0x2A)
GRAY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---- 기본 스타일 / 여백 ----
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.35

for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)

def _run(p, text, size=10.5, bold=False, color=None, italic=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r

def shade(el, hexcolor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    el.get_or_add_tcPr().append(sh) if el.tag.endswith("tc") else el.get_or_add_pPr().append(sh)

def bar(text, sub=None):
    """남색 바 제목"""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c._tc, "1F3364")
    c.width = Cm(16.2)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    if sub:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        _run(p2, sub, size=9.5, color=RGBColor(0xC8, 0xD4, 0xE8))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=11.5, bold=True, color=NAVY)

def body(text, indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    _run(p, text)
    return p

def bullet(text, indent=0.5, mark="•"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent + 0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(3)
    _run(p, f"{mark} {text}")
    return p

def rich_bullet(segs, indent=0.5, mark="•"):
    """segs = [(text, bold, color), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent + 0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(3)
    _run(p, f"{mark} ")
    for text, b, col in segs:
        _run(p, text, bold=b, color=col)
    return p

def flowbox(title, text, fill="EEF2F8", border="1F3364"):
    """강조 박스"""
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c._tc, fill)
    tcPr = c._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "12" if edge == "left" else "4")
        e.set(qn("w:color"), border)
        borders.append(e)
    tcPr.append(borders)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    if title:
        _run(p, title, size=10.5, bold=True, color=NAVY)
        p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    _run(p, text, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c._tc, "1F3364")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        _run(p, htxt, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _run(p, v, size=9.5, bold=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run(p, v, size=9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# =====================================================================
# 표지 영역
# =====================================================================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
_run(p, "M1013 VLA 모방학습 · 전이학습 연구방향 정리", size=18, bold=True, color=NAVY)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
_run(p, "이종 형상(Cross-embodiment) 로봇 간 모방학습 전이 — 축 불일치 문제의 재정의", size=11, color=GRAY)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_after = Pt(16)
_run(p, "2026. 09. 09.", size=10, color=GRAY)

# =====================================================================
# 연구 주제 한 문장
# =====================================================================
bar("연구 주제 (한 문장)")

flowbox(
    "■ 대표안",
    "저자유도(5축) 저가 매니퓰레이터에서 수집한 모방학습 데모가 고자유도(6축) 산업용 매니퓰레이터의 "
    "VLA 정책 학습에 기여하는 정도를, 엔드이펙터 pose 기반 공통 행동공간에서 정량화한다.",
    fill="EEF2F8",
)

body("· 보조안(실용 강조): 서로 다른 자유도의 로봇 데모를 EE pose 공통 표현으로 통합했을 때, "
     "목표 로봇의 실물 데모를 몇 개까지 대체할 수 있는지를 데이터 효율 곡선으로 규명한다.")
body("· 보조안(학술 강조): 형상 격차(morphology gap)와 도메인 격차(sim-to-real gap)가 "
     "VLA 전이 성능에 미치는 영향을 분리 측정하고, 자유도 결손 축을 보상하는 residual 어댑터 구조를 제안한다.")

doc.add_paragraph()

# =====================================================================
# 1. 현황 진단
# =====================================================================
bar("1. 현황 진단 — 문제의 재정의")

h2("1.1  기존 인식")
body("OMX(그리퍼 포함 6축)와 M1013(그리퍼 포함 7축)의 축 개수가 달라 모방학습 데이터를 "
     "그대로 전이할 수 없다는 것이 기존 문제 인식이었음.")

h2("1.2  데이터 확인 결과")
table(
    ["구분", "팔 자유도", "그리퍼", "행동벡터", "비고"],
    [
        ["omx_f (소스)", "5축", "1", "7D (EE pose)", "joint1 yaw / joint2·3·4 pitch / joint5 roll"],
        ["M1013 (목표)", "6축", "1", "7D (EE pose)", "2026-09-20 도착 예정"],
    ],
    widths=[3.0, 2.2, 1.8, 3.2, 6.0],
)

body("현재 사전학습 데이터셋(robot_type = omx_f, 159 에피소드 / 76,345 프레임)은 이미 "
     "[x, y, z, rx, ry, rz, gripper] 7차원 task-space 표현으로 구축되어 있음.")

flowbox(
    "■ 결론 1 — 축 개수는 이미 해결된 문제",
    "FK로 EE pose를 추출한 시점에서 관절 개수는 표현에서 사라졌음. OMX도 7D, M1013도 7D로 동일하며, "
    "정책 출력(EE pose)을 IK로 목표 로봇의 관절값으로 되돌리는 경로도 이미 확보되어 있음.",
    fill="EEF2F8",
)

flowbox(
    "■ 결론 2 — 실질 병목은 '자세 커버리지 결손'",
    "omx_f는 5축이므로 EE yaw를 독립적으로 제어할 수 없음. joint1이 베이스 회전축이라 "
    "EE의 yaw가 위치에 종속(yaw ≈ atan2(y, x))되며, 결과적으로 7차원 행동벡터 중 실질 자유도는 6개, "
    "SE(3) 내에서 5차원 부분다양체만 커버함. 즉 사전학습 데이터에 'yaw를 자유롭게 사용하는 조작' 사례가 "
    "원천적으로 존재하지 않는 것이 진짜 격차임.",
    fill="FBF0EC", border="C0502A",
)

body("이 진단은 문제인 동시에 연구 프레이밍의 근거가 됨 — 자유도 결손을 제거 대상이 아니라 "
     "통제 변수로 삼으면 그 자체가 실험 설계가 됨.")

doc.add_page_break()

# =====================================================================
# 방향 1
# =====================================================================
bar("방향 1. OMX → M1013 직접 전이 효과의 정량화 (권장)",
    "'모방학습 데이터가 값어치를 하는가'에 대한 직접적 답")

h2("핵심 아이디어")
body("지도교수님 요구사항의 본질은 '모방학습 데이터가 실제로 효과가 있는가'를 보이는 것임. "
     "이를 보이는 표준 형식은 데이터 효율 곡선(data efficiency curve)이며, 여기서 도출되는 "
     "등가 데이터량(data equivalence) 지표가 결론 문장이 됨. 보유 중인 omx_f 데이터 159 에피소드만으로 "
     "추가 장비 없이 즉시 착수 가능함.")

h2("1-A.  데이터 효율 곡선")
bullet("x축: M1013 실물 데모 개수 (10 / 30 / 60 / 120)")
bullet("y축: 태스크 성공률")
bullet("곡선: ① scratch(M1013 데모만)  ② 비전 인코더만 전이  ③ 전체(비전+행동) 전이")
rich_bullet([("도출 지표: ", True, None),
             ("\"OMX 데모 159개 = M1013 실물 데모 N개 값어치\"", True, ACC),
             (" — 이 한 문장이 보고의 결론", False, None)])

h2("1-B.  전이 층위 3단 분해")
body("어디까지 전이되고 어디부터 전이되지 않는지를 분해하여, 형상 격차의 영향 지점을 특정함.")

table(
    ["층위", "전이 범위", "행동공간 처리", "확인 사항"],
    [
        ["L1 (지각)", "비전 인코더 + 언어 조건",
         "action head는 신규 학습", "축 문제를 완전히 회피한 전이 하한선"],
        ["L2 (전체)", "비전 + task-space 행동",
         "7D 행동 그대로 전이", "yaw 결손 부분공간이 도움인가 방해인가"],
        ["L3 (어댑터)", "L2 + yaw residual head",
         "결손 1축만 M1013 데모로 학습", "부족 자유도만 선택적으로 보강 가능한가"],
    ],
    widths=[2.4, 4.0, 4.2, 5.6],
)

flowbox(
    "■ 이 설계의 강점 — 실패해도 결과가 남음",
    "L1이 L2보다 우수하게 나올 경우, 이는 실패가 아니라 '형상 격차가 행동공간 전이를 오염시킨다'는 "
    "독립적인 연구 결과임. 어떤 순위가 나오더라도 해석 가능한 결론이 도출되는 구조.",
    fill="FBF0EC", border="C0502A",
)

h2("1-C.  평가 태스크")
bullet("기존 pick & place 태스크를 기본 평가군으로 유지 (omx_f 데이터와 동일 태스크)")
bullet("yaw 자유도가 필요한 태스크(비스듬히 놓인 물체 파지 등)를 별도 평가군으로 추가")
bullet("→ 두 평가군의 성능 격차가 곧 '자유도 결손의 비용'을 직접 보여주는 지표가 됨")

doc.add_page_break()

# =====================================================================
# 방향 2 — OMX를 사용하지 않는 경로
# =====================================================================
bar("방향 2. OMX를 사용하지 않는 사전학습 경로",
    "자유도 결손 자체를 원천 회피하는 대안")

h2("핵심 아이디어")
body("방향 1이 'omx_f의 자유도 결손을 통제 변수로 삼는' 접근이라면, 방향 2는 사전학습 소스를 "
     "omx_f 이외의 것으로 대체하여 결손 자체를 발생시키지 않는 접근임. "
     "형상 격차가 줄어드는 대신 다른 종류의 격차(도메인·시점·구현체)가 새로 생기므로, "
     "무엇을 감수할지의 선택 문제가 됨.")

h2("2-A.  Isaac Sim 내 M1013 자체 데모  (최우선)")
rich_bullet([("형상 격차 = 0. ", True, ACC),
             ("목표 로봇과 동일한 기구학으로 대량 데모를 생성하므로 자유도 결손·IK 재타겟팅 문제가 "
              "모두 소멸함. 남는 격차는 sim-to-real 하나뿐.", False, None)])
bullet("환경 구축 완료 상태 활용 가능 — 상판 높이 0.4624, TCP(파지점) 0.0725, 그리퍼 조인트 5종 + 플랜지 검증 완료")
bullet("도메인 랜덤화(조명·텍스처·물체 위치·카메라 외부파라미터)로 sim-to-real 격차 완화")
bullet("제약: 접촉 물리(파지 순간의 미끄러짐·변형)의 재현도가 실물 대비 낮음")

h2("2-B.  동형 6축 실물 로봇")
table(
    ["후보", "자유도", "형상 격차", "비고"],
    [
        ["Doosan A0509", "6축", "매우 작음", "동일 제조사·동일 kinematic family. URDF 교체 수준으로 대응 가능"],
        ["xArm 6", "6축", "작음", "구조 동일(6축+그리퍼). 저가 실물 대안"],
        ["UR3e / UR5e", "6축", "작음", "구조 동일. 공개 데이터·레퍼런스가 풍부"],
    ],
    widths=[3.4, 1.8, 2.2, 8.8],
)
bullet("자유도가 일치하므로 yaw 결손이 발생하지 않고, EE pose 전이가 온전히 성립함")
bullet("제약: 장비 확보 가능 여부가 관건. 확보 시 방향 1보다 전이 상한이 높을 것으로 예상")

h2("2-C.  공개 대규모 로봇 데이터셋")
bullet("Open X-Embodiment / DROID(Franka 7축) / BridgeData V2(WidowX 6DOF)")
bullet("자체 수집 없이 '대규모 사전학습' 항목을 실험에 추가 가능")
rich_bullet([("자유도가 제각각인 로봇들을 이미 통합한 데이터셋이므로, ", False, None),
             ("축 불일치 처리의 표준 레시피(최대 DOF 패딩 + 마스킹 + embodiment 토큰)", True, None),
             ("를 그대로 차용할 수 있음 — 본 연구의 축 문제에 대한 기성 해법 참조원", False, None)])
bullet("제약: 태스크·작업공간·카메라 시점이 본 연구 세팅과 달라 그대로는 성능 이득이 제한적")

h2("2-D.  인간 시연 영상 기반 사전학습")
bullet("Ego4D / Something-Something 등 인간 조작 영상으로 시각 표현만 사전학습 (R3M·VC-1 계열)")
bullet("로봇 행동 레이블이 없으므로 자유도 개념 자체가 존재하지 않음 → 축 불일치 원천 무관")
bullet("방향 1의 L1(지각 전이)과 결합 가능 — 비전 인코더 초기값으로 사용하고 행동은 M1013 데모로 학습")
bullet("제약: 행동 사전지식이 전이되지 않아 단독으로는 데이터 효율 이득이 작음")

flowbox(
    "■ 방향 1과 방향 2를 함께 수행할 경우",
    "실물 OMX(형상 다름 · 실제 데이터)와 시뮬 M1013(형상 동일 · 시뮬 데이터)을 나란히 비교하면 "
    "「형상 격차와 도메인 격차 중 무엇이 VLA 전이에 더 해로운가」라는 독립적인 연구 질문이 성립함. "
    "두 축을 분리 측정한 선행 사례가 많지 않아 기여도가 높음.",
)

doc.add_page_break()

# =====================================================================
# 3. 연구의 한계점
# =====================================================================
bar("3. 연구의 한계점", "EE pose 기반 행동공간이 필연적으로 안고 가는 제약")

h2("3.1  IK 특이점(singularity) 문제")
body("본 연구는 정책이 EE pose를 출력하고 이를 IK로 목표 로봇의 관절값으로 변환하는 구조를 전제함. "
     "이 구조는 축 불일치를 해소하는 대신, 관절공간 제어에는 없는 특이점 문제를 새로 떠안게 됨.")

table(
    ["특이점 유형", "발생 조건", "증상"],
    [
        ["손목 특이점", "손목 축 2개가 정렬(4·6축 동축)",
         "해당 회전 자유도 상실, 관절속도 급증"],
        ["어깨 특이점", "손목 중심이 1축 회전축 선상에 위치",
         "베이스 회전 관절의 속도 발산"],
        ["팔꿈치 특이점", "팔이 완전히 펴진 작업공간 경계",
         "IK 해 소멸 또는 해의 불연속 점프"],
    ],
    widths=[3.0, 6.4, 6.8],
)

bullet("정책이 출력한 pose가 특이점 근방일 경우, 자코비안 역행렬이 발산하여 관절 지령속도가 폭주함")
bullet("연속된 프레임에서 IK 해가 다른 분지(branch)로 점프하면, 궤적이 매끄럽던 학습 데이터와 달리 "
       "실행 시 급격한 자세 변화가 발생함")
rich_bullet([("사전학습 소스(omx_f 5축)와 목표 로봇(M1013 6축)의 특이점 분포가 서로 다르므로, "
              "OMX에서는 안전했던 궤적이 M1013에서는 특이점을 통과할 수 있음", True, ACC)])

h2("3.2  특이점 회피를 위한 작업공간 제약")
body("특이점 통과를 실물에서 허용할 수 없으므로, 작업대의 위치·높이·물체 배치 범위를 "
     "특이점에서 충분히 떨어진 영역으로 제한하여 실험을 구성하게 됨. "
     "이는 안전상 불가피한 선택이지만, 동시에 본 연구 결과의 적용 범위를 좁히는 한계로 작용함.")

bullet("작업대를 로봇 전방 특정 거리·높이 대역으로 고정 → 학습·평가 작업공간이 전체 도달영역의 일부로 한정됨")
bullet("정책이 학습한 분포가 제한된 영역에 한정되므로, 작업대 위치가 바뀌면 성능 보장이 어려움 "
       "(일반화 범위의 제약)")
bullet("omx_f의 작업공간(리치 약 30cm)과 M1013(약 1.3m급)의 스케일 차이가 커, "
       "정규화 후에도 두 로봇이 공통으로 안전한 영역은 더욱 좁아짐")
rich_bullet([("결과 해석 시 명시 필요: ", True, None),
             ("보고되는 성공률은 '특이점 회피가 보장된 제한 작업영역'에서의 값이며, "
              "전 작업공간에 대한 성능이 아님", False, None)])

h2("3.3  완화 방안 (한계 인정 + 대응)")
table(
    ["구분", "방안", "효과"],
    [
        ["작업영역 선정", "manipulability index 기반으로 특이점에서 먼 영역을 정량 산출하여 작업대 배치",
         "제약을 임의 설정이 아닌 근거 있는 설계로 전환"],
        ["IK 수치안정화", "damped least squares(DLS) 기반 IK로 특이점 근방 속도 발산 억제",
         "폭주 방지, 단 해의 정확도는 일부 희생"],
        ["해 연속성", "직전 프레임 관절값을 시드로 사용해 IK 분지 점프 억제",
         "궤적 불연속 완화"],
        ["학습 단계 반영", "특이점 근접도를 학습 손실의 페널티 항으로 추가",
         "정책이 특이점 근방 pose를 애초에 회피하도록 유도"],
        ["구조적 대안", "방향 2-A(Isaac Sim 내 M1013 자체 데모) 채택",
         "재타겟팅 자체가 없어 특이점 문제의 상당 부분이 소멸"],
    ],
    widths=[2.8, 7.2, 6.2],
)

flowbox(
    "■ 한계점의 연구적 의미",
    "IK 특이점과 그로 인한 작업공간 제약은 EE pose 공통 표현을 채택한 대가임. "
    "따라서 본 연구의 결론은 '축 불일치를 해소하면 전이가 가능하다'가 아니라 "
    "'축 불일치를 EE pose 표현으로 해소할 때, 특이점 회피를 위한 작업영역 제한이라는 비용이 발생하며 "
    "그 비용을 감안해도 전이 이득이 존재하는가'라는 형태로 기술하는 것이 정확함.",
    fill="FBF0EC", border="C0502A",
)

doc.add_page_break()

# =====================================================================
# 4. 실무 체크리스트
# =====================================================================
bar("4. 실무 체크리스트 — 로봇 도착 전 준비사항")

h2("4.1  M1013 실물 데모 수집 시")
rich_bullet([("yaw 다양성을 의도적으로 포함할 것. ", True, ACC),
             ("OMX 데이터에 존재하지 않는 축이므로, 파인튜닝 데이터가 이를 채워주지 않으면 "
              "전이 실패가 아니라 단순 데이터 부재로 인한 실패가 됨.", False, None)])
bullet("평가군 분리를 전제로 수집 — yaw 불필요 태스크 / yaw 필요 태스크를 처음부터 구분하여 기록")

h2("4.2  작업대 배치 및 특이점 회피")
bullet("작업대 위치 확정 전 manipulability index 맵을 산출하여 안전 영역을 먼저 정의")
bullet("확정된 작업영역을 문서화하고, 모든 데모·평가를 동일 영역 내에서 수행")
bullet("IK 솔버에 DLS 감쇠계수 및 직전 해 시드 적용 여부를 사전 확정")

h2("4.3  작업공간 정규화")
bullet("OMX 리치 약 30cm, M1013 약 1.3m급으로 스케일 차이가 큼")
bullet("작업공간 대비 비율 좌표로 정규화하면 전이 안정성이 크게 향상됨")

h2("4.4  기하 정합 유지")
bullet("TCP = 파지점 0.0725 기준 유지")
bullet("Isaac 상판 높이 0.4624 기준 유지")
bullet("그리퍼 조인트 5종 + 로봇 플랜지 연결 검증 완료 상태 유지")

doc.add_paragraph()

# =====================================================================
# 5. 보고 요약
# =====================================================================
bar("5. 지도교수님 보고용 요약")

flowbox(
    "■ 한 문단 요약",
    "축 개수 차이는 엔드이펙터 pose 기반 표현으로 이미 해소되었으며, 남은 실질 격차는 "
    "5축 팔의 yaw 자유도 결손입니다. 이를 제거 대상이 아니라 통제 변수로 삼아, "
    "OMX 사전학습 유무에 따른 M1013의 데이터 효율 곡선을 비교하여 모방학습 데이터의 값어치를 "
    "'실물 데모 N개 등가'라는 정량 지표로 제시하겠습니다. 다만 EE pose 표현은 IK 특이점을 "
    "수반하므로, 특이점 회피를 위한 제한 작업영역 내의 결과임을 함께 명시할 예정입니다. "
    "OMX를 쓰지 않는 대안으로는 Isaac Sim 내 M1013 자체 데모가 가장 유력합니다.",
    fill="EEF2F8",
)

h2("진행 계획")
table(
    ["단계", "내용", "시점"],
    [
        ["1", "manipulability 맵 산출 → 작업대 배치 및 안전 작업영역 확정", "로봇 도착 전"],
        ["2", "사전학습 2종(scratch / omx_f) 파이프라인 구축 및 IK 솔버 안정화", "로봇 도착 전"],
        ["3", "M1013 실물 데모 수집 (yaw 다양성 포함, 평가군 분리)", "09-20 이후"],
        ["4", "데이터 효율 곡선 측정 및 등가 데이터량 산출", "수집 완료 후"],
        ["5", "L1/L2/L3 전이 층위 분해 실험", "4단계와 병행"],
        ["6", "(선택) Isaac Sim 내 M1013 데모와의 비교 — 형상 격차 vs 도메인 격차", "5단계 이후"],
    ],
    widths=[1.6, 10.6, 4.0],
)

out = "/home/kim/m1013/대화록 및 PPT/M1013_연구방향_20260909.docx"
doc.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
