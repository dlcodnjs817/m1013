#!/usr/bin/env python3
"""2026-08-24 업무 정리 PPT — STEP 5-1/5-2 카메라 구도 정합.

  python3 make_report_pptx_0824.py
  → /home/kim/m1013/대화록 및 PPT/M1013_카메라정합_20260824.pptx

디자인: 남색 바 + flowbox 기존 체계 (0814/0820/0821 스크립트 계승)
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

NAVY = RGBColor(0x2F, 0x3D, 0x9E)
INK = RGBColor(0x15, 0x18, 0x1D)
GREY = RGBColor(0x62, 0x6B, 0x78)
RED = RGBColor(0xA9, 0x33, 0x1D)
GREEN = RGBColor(0x1F, 0x6B, 0x4A)
ORANGE = RGBColor(0xB4, 0x5F, 0x06)
BGSOFT = RGBColor(0xED, 0xEF, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LILAC = RGBColor(0xB9, 0xC2, 0xF0)
LINE = RGBColor(0xC6, 0xCC, 0xD8)

FONT = "맑은 고딕"
MONO = "Consolas"
SW, SH = Cm(33.867), Cm(19.05)
BASE = "/home/kim/m1013"
OUTDIR = os.path.join(BASE, "대화록 및 PPT")
IMG = os.path.join(OUTDIR, "img_0824")
SIMOUT = os.path.join(BASE, "sim_out")

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def _set(run, size=15, bold=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def slide():
    return prs.slides.add_slide(BLANK)


def bar(s, title, step=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Cm(2.0))
    r.fill.solid(); r.fill.fore_color.rgb = NAVY
    r.line.fill.background()
    tf = r.text_frame
    tf.margin_left = Cm(0.9); tf.margin_top = Cm(0.26)
    p = tf.paragraphs[0]
    if step:
        run = p.add_run(); run.text = step + "  "
        _set(run, 15, True, LILAC)
    run = p.add_run(); run.text = title
    _set(run, 21, True, WHITE)
    return r


def textbox(s, x, y, w, h):
    tb = s.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    return tb.text_frame


def bullets(tf, items, size=15):
    first = True
    for it in items:
        lv, txt = it[0], it[1]
        opt = it[2] if len(it) > 2 else {}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lv
        p.space_after = Pt(opt.get("after", 6))
        marks = {0: "▪  ", 1: "–  ", 2: "·  "}
        run = p.add_run()
        run.text = ("" if opt.get("nomark") else marks.get(lv, "")) + txt
        _set(run, opt.get("size", size - lv), opt.get("bold", False),
             opt.get("color", INK), MONO if opt.get("mono") else FONT)


def table(s, x, y, w, rows, widths=None, size=12.5, header=True, rh=Cm(0.8)):
    shp = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, rh * len(rows))
    t = shp.table
    if widths:
        total = sum(widths)
        for i, cw in enumerate(widths):
            t.columns[i].width = int(w * cw / total)
    for ri, row in enumerate(rows):
        t.rows[ri].height = rh
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci)
            c.margin_left = Cm(0.16); c.margin_right = Cm(0.16)
            c.margin_top = Cm(0.05); c.margin_bottom = Cm(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            txt = cell if isinstance(cell, str) else cell[0]
            opt = {} if isinstance(cell, str) else cell[1]
            c.fill.solid()
            c.fill.fore_color.rgb = (NAVY if (header and ri == 0)
                                     else opt.get("bg", WHITE if ri % 2 else BGSOFT))
            p = c.text_frame.paragraphs[0]
            run = p.add_run(); run.text = txt
            _set(run, opt.get("size", size),
                 opt.get("bold", header and ri == 0),
                 WHITE if (header and ri == 0) else opt.get("color", INK),
                 MONO if opt.get("mono") else FONT)
    return t


def flowbox(s, x, y, w, h, txt, fill=NAVY, fg=WHITE, size=12, bold=True):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    b.fill.solid(); b.fill.fore_color.rgb = fill
    b.line.color.rgb = LINE; b.line.width = Pt(0.75)
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = Cm(0.1); tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.06); tf.margin_bottom = Cm(0.06)
    for i, line in enumerate(txt.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = line
        _set(run, size if i == 0 else size - 2, bold if i == 0 else False, fg)
    return b


def arrow(s, x, y, w=Cm(0.7)):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Cm(0.5))
    a.fill.solid(); a.fill.fore_color.rgb = GREY
    a.line.fill.background()


def foot(s, txt="M1013 전이 · STEP 5 카메라 정합 · 2026-08-24"):
    tf = textbox(s, Cm(0.9), SH - Cm(0.72), SW - Cm(1.8), Cm(0.6))
    p = tf.paragraphs[0]
    run = p.add_run(); run.text = txt
    _set(run, 10, False, GREY)


def caption(s, x, y, w, txt, size=11.5, bold=False, color=GREY):
    tf = textbox(s, x, y, w, Cm(1.2))
    p = tf.paragraphs[0]
    run = p.add_run(); run.text = txt
    _set(run, size, bold, color)


def pic(s, path, x, y, w):
    if os.path.exists(path):
        return s.shapes.add_picture(path, x, y, width=w)
    caption(s, x, y, w, f"[누락: {os.path.basename(path)}]", 12, True, RED)


# ═══════════ 1. 표지 ═══════════
s = slide()
r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
r.fill.solid(); r.fill.fore_color.rgb = NAVY; r.line.fill.background()
tf = textbox(s, Cm(2.2), Cm(4.6), Cm(29.5), Cm(8))
p = tf.paragraphs[0]
run = p.add_run(); run.text = "STEP 5-1 · 5-2 — 카메라 구도 정합"
_set(run, 36, True, WHITE)
p = tf.add_paragraph(); p.space_before = Pt(10)
run = p.add_run()
run.text = "OMX 학습 영상(손목캠+정면캠)과 같은 구도의 카메라 2대를 Isaac M1013 씬에 장착"
_set(run, 19, False, LILAC)
p = tf.add_paragraph(); p.space_before = Pt(24)
run = p.add_run()
run.text = "핵심 산출물:  cameras_v6match.json (확정 파라미터) · cam_match_v1.png (실물 vs sim 비교시트)"
_set(run, 14, True, WHITE)
p = tf.add_paragraph(); p.space_before = Pt(6)
run = p.add_run()
run.text = "다음 단계:  STEP 5-3 closed-loop — 이 카메라의 렌더를 관측으로 정책 실시간 제어"
_set(run, 14, True, WHITE)
tf2 = textbox(s, Cm(2.2), SH - Cm(1.7), Cm(24), Cm(0.9))
run = tf2.paragraphs[0].add_run()
run.text = "2026-08-24 (월) · 이채원"
_set(run, 13, False, LILAC)

# ═══════════ 2. 전체 로드맵에서의 위치 ═══════════
s = slide()
bar(s, "전체 로드맵에서 오늘의 위치", "00 파이프라인")
steps = [
    ("STEP 1\n에셋·IK·배치", GREEN, "8/20 완료"),
    ("STEP 2\nv6-M1013 데이터셋", GREEN, "8/20 완료"),
    ("STEP 3\nIsaac 물리 검증", GREEN, "8/20 완료"),
    ("STEP 4\nACT 학습·평가", GREEN, "8/20 완료"),
    ("STEP 5\n카메라·closed-loop", ORANGE, "◀ 오늘 (진행 중)"),
    ("STEP 6\n실기 전이", GREY, "로봇 도착 후"),
]
bw, bh, gap = Cm(4.7), Cm(2.1), Cm(0.75)
x0 = Cm(0.9); y0 = Cm(2.9)
for i, (txt, col, st) in enumerate(steps):
    x = x0 + i * (bw + gap)
    flowbox(s, x, y0, bw, bh, txt, fill=col, size=13)
    caption(s, x, y0 + bh + Cm(0.1), bw, st, 11, True,
            GREEN if col == GREEN else (ORANGE if col == ORANGE else GREY))
    if i < len(steps) - 1:
        arrow(s, x + bw + Cm(0.03), y0 + bh / 2 - Cm(0.25))
# STEP5 내부 분해
tf = textbox(s, Cm(0.9), Cm(6.2), Cm(32), Cm(0.8))
run = tf.paragraphs[0].add_run(); run.text = "STEP 5 내부 분해 — 오늘은 5-1 · 5-2 완료"
_set(run, 15, True, INK)
sub = [
    ("5-1  씬에 카메라 추가\n정면캠(월드 고정)\n손목캠(link_6 부착)", GREEN, "완료 ✔"),
    ("5-2  구도 정합\n원본 영상과 위치·각도·FOV\n일치 + 비교시트", GREEN, "완료 ✔"),
    ("5-3  closed-loop 가동\n렌더 관측 → 정책 → 제어\n시각 도메인 갭 측정", ORANGE, "다음 시작점 ▶"),
    ("5-4  (조건부) co-training\n159ep 재렌더 A/B\n갭이 크면 채택", GREY, "5-3 결과에 따라"),
]
bw2 = Cm(7.5)
for i, (txt, col, st) in enumerate(sub):
    x = x0 + i * (bw2 + gap)
    flowbox(s, x, Cm(7.1), bw2, Cm(2.6), txt, fill=col, size=12.5)
    caption(s, x, Cm(9.8), bw2, st, 11.5, True,
            GREEN if col == GREEN else (ORANGE if col == ORANGE else GREY))
    if i < len(sub) - 1:
        arrow(s, x + bw2 + Cm(0.03), Cm(7.1) + Cm(1.05))
tf = textbox(s, Cm(0.9), Cm(11.2), Cm(32), Cm(6.4))
bullets(tf, [
    (0, "왜 카메라 정합이 먼저인가: ACT 정책의 관측 = 두 카메라 영상. closed-loop 은 sim 렌더가 그 자리를 대신하므로,", {"size": 14.5}),
    (1, "학습 때 본 화면과 구도가 다르면 '시각 도메인 갭'과 '구도 불일치'가 섞여서 실패 원인을 분리할 수 없음", {"size": 13.5}),
    (1, "구도를 먼저 못박아 두면 남는 차이 = 순수 도메인 갭(로봇 외형·배경·질감) → 5-3 에서 정량 측정 가능", {"size": 13.5, "bold": True}),
    (0, "제약: 원본은 OMX(소형)·sim 은 M1013(대형) — '로봇 기준'이 아닌 작업 공간(큐브·그리퍼·테이블) 기준 정합을 선택", {"size": 14.5}),
])
foot(s)

# ═══════════ 3. 오늘 작업 파이프라인 ═══════════
s = slide()
bar(s, "오늘 작업 파이프라인 — 5단계", "00 파이프라인")
flow = [
    ("① 원본 영상 확보\n호스트 심링크 깨짐\n→ 컨테이너 HF 캐시", NAVY),
    ("② 카메라 구성 분석\ncamera1=손목캠\ncamera2=정면캠 판별", NAVY),
    ("③ 정면캠 탐색\n3라운드 × 후보 3~4개\n위치·FOV 수렴", NAVY),
    ("④ 손목캠 탐색\n부착·롤 실증 3라운드\n장착 오프셋 확정", NAVY),
    ("⑤ 확정·산출\ncameras_v6match.json\ncam_match_v1.png", GREEN),
]
bw, bh, gap = Cm(5.85), Cm(2.5), Cm(0.7)
x0 = Cm(0.9); y0 = Cm(2.7)
for i, (txt, col) in enumerate(flow):
    x = x0 + i * (bw + gap)
    flowbox(s, x, y0, bw, bh, txt, fill=col, size=12.5)
    if i < len(flow) - 1:
        arrow(s, x + bw + Cm(0.0), y0 + bh / 2 - Cm(0.25))
# 반복 루프 구조
tf = textbox(s, Cm(0.9), Cm(5.7), Cm(32), Cm(0.8))
run = tf.paragraphs[0].add_run(); run.text = "③④ 내부의 반복 단위 — Isaac 부팅이 느려서(회당 ~1분) 이렇게 설계"
_set(run, 15, True, INK)
loop = [
    ("실측 역산으로\n초기값·범위 산출", BGSOFT, INK),
    ("Isaac 1회 부팅에\n후보 3~4개 동시 렌더", BGSOFT, INK),
    ("원본과 픽셀 정량 비교\n(주사위 위치·크기 %)", BGSOFT, INK),
    ("차이의 원인 해석\n→ 다음 후보 설계", BGSOFT, INK),
]
bw2 = Cm(7.3)
for i, (txt, bg, fg) in enumerate(loop):
    x = x0 + i * (bw2 + gap)
    flowbox(s, x, Cm(6.6), bw2, Cm(1.9), txt, fill=bg, fg=fg, size=12)
    if i < len(loop) - 1:
        arrow(s, x + bw2 + Cm(0.0), Cm(6.6) + Cm(0.7))
caption(s, x0, Cm(8.6), Cm(32), "↺ 3라운드 반복 후 수렴 — 정면캠(A~D → E~H → F1~F3), 손목캠(W1~4 → 롤 스윕 → V1~3)", 12, True, GREY)
rows = [
    ["단계", "입력", "출력", "사용 도구"],
    ["①", "v6 심링크(깨짐)·컨테이너", ("omx v4 원본 mp4 (159ep×2뷰)", {"mono": True}), "docker exec + 컨테이너 ffmpeg"],
    ["②", "원본 mp4", "카메라 역할·스펙·이상 에피소드 판별", "프레임 그리드 추출·육안+픽셀 분석"],
    ["③", "기준 프레임(정면 홈·파지)", ("F1 확정: pos·look_at·f15mm", {"mono": True}), ("frontcam_probe.py", {"mono": True})],
    ["④", "기준 프레임(손목 홈·파지)", ("V3 확정: local_pos·롤·f18mm", {"mono": True}), ("wristcam_probe.py", {"mono": True})],
    ["⑤", "③④ 확정값", ("cameras_v6match.json · cam_match_v1.png", {"mono": True}), "JSON 직렬화·비교시트 합성"],
]
table(s, Cm(0.9), Cm(9.5), Cm(32.1), rows, widths=[1.2, 6.4, 8.2, 6.8], size=12, rh=Cm(1.25))
foot(s)

# ═══════════ 4. 원본 데이터셋 분석 ═══════════
s = slide()
bar(s, "원본(OMX v6) 카메라 구성 — 손목캠 + 정면캠", "01 분석")
iw = Cm(10.8); ih = iw * 3 / 4
pic(s, os.path.join(IMG, "ref_c2_home.png"), Cm(0.9), Cm(2.5), iw)
pic(s, os.path.join(IMG, "wrist_ref_grid.png"), Cm(12.2), Cm(2.5), iw)
caption(s, Cm(0.9), Cm(2.5) + ih + Cm(0.08), iw, "camera2 = 정면 고정캠 — 낮은 앵글 · pick 좌 · place 우", 11.5, True, INK)
caption(s, Cm(12.2), Cm(2.5) + ih + Cm(0.08), iw, "camera1 = 손목캠 — 주사위 위→아래 진입, 파지 시 하단", 11.5, True, INK)
rows = [
    ["항목", "값"],
    ["해상도·FPS", "640×480 · 30fps (두 뷰 동일) · 에피소드당 ~460프레임(≈15초)"],
    ["camera1 (손목캠)", "그리퍼 장착 하향 시점 · 이미지 위 = 접근 방향 · 파지점은 화면 하단 중앙 · 핑거가 모서리에 걸림"],
    ["camera2 (정면캠)", "테이블 상판 높이+약 10cm · 주사위와 25~30cm (픽셀 크기로 역산) · HFOV 60~70° 추정 · 배경에 사람·파티션"],
    ["에피소드 변주", "place 쪽 주사위 색·개수가 에피소드마다 다름 (흰1/노1/초·주·빨3 …) — 자연 도메인 랜덤화"],
]
table(s, Cm(0.9), Cm(11.7), Cm(32.1), rows, widths=[4.2, 17.8], size=12, rh=Cm(1.1))
tf = textbox(s, Cm(0.9), Cm(17.4), Cm(32), Cm(1.4))
bullets(tf, [
    (0, "부수 발견: v4 ep0 은 카메라 매핑 불량(정면↔먹통 회색) — v6 는 심링크 리매핑(v6 ep0→v4 ep1)으로 이미 제외되어 있었음", {"size": 12.5, "color": ORANGE, "bold": True}),
])
foot(s)

# ═══════════ 5. 트러블슈팅 ①② 데이터 접근 ═══════════
s = slide()
bar(s, "트러블슈팅 ①② — 원본 영상이 호스트에 없다", "02 트러블슈팅")
rows = [
    ["#", "증상", "원인", "해결"],
    ["①", "v6 videos/ 의 mp4 가 전부 열리지 않음",
     "전부 심링크인데 목적지 폴더(modeldataset/datasets/omx_…_v4_162)가 삭제됨",
     ("컨테이너 physical_ai_server 의 HF 캐시에서 원본 발견 → docker cp 로 추출", {"bold": True})],
    ["②", "호스트에 ffmpeg 없음 + cv2 는 NumPy 2.x 충돌",
     "호스트 파이썬 환경이 렌더/ML 용으로만 구성",
     ("컨테이너 안 ffmpeg 로 프레임 추출 후 꺼내오는 우회 (docker exec + docker cp)", {"bold": True})],
]
table(s, Cm(0.9), Cm(2.6), Cm(32.1), rows, widths=[0.9, 6.5, 7.3, 8.3], size=12, rh=Cm(1.5))
tf = textbox(s, Cm(0.9), Cm(7.6), Cm(32), Cm(3.2))
bullets(tf, [
    (0, "원본 실체 경로 (이후에도 필요 — 재렌더·co-training 시):", {"size": 14, "bold": True}),
    (1, "physical_ai_server:/root/.cache/huggingface/lerobot/dlcodnjs/omx_act_pick_and_place_v4_162/videos/", {"size": 12, "mono": True}),
    (0, "교훈: 깨진 심링크는 '데이터 소실'이 아니라 '참조 대상 이동'일 수 있음 — readlink -f 로 목적지부터 확인", {"size": 13.5}),
])
iw2 = Cm(16.0)
pic(s, os.path.join(IMG, "c2_grid.png"), Cm(0.9), Cm(11.0), iw2)
caption(s, Cm(0.9), Cm(11.0) + iw2 * 480 / 1280 + Cm(0.08), iw2,
        "v4 ep0 의 camera2 — 에피소드 내내 동일한 회색 정지화면 (매핑 불량의 증거, v6 에서 제외됨)", 11.5, True, INK)
tf = textbox(s, Cm(17.8), Cm(11.2), Cm(15.2), Cm(5.6))
bullets(tf, [
    (0, "ep0 이상 판별 과정:", {"size": 13.5, "bold": True}),
    (1, "f0/f240/f420 세 프레임이 픽셀 동일 → 고정 먹통 캠", {"size": 12.5}),
    (1, "ep10~150 은 camera1 이 움직이는 손목캠으로 정상", {"size": 12.5}),
    (1, "v6 심링크가 ep 번호를 리매핑(ep0→v4 ep1)함을 확인 → 학습 데이터는 무결", {"size": 12.5, "color": GREEN, "bold": True}),
])
foot(s)

# ═══════════ 6. 방법 선택 — 채택 vs 대안 ═══════════
s = slide()
bar(s, "카메라 정합 방법 — 채택한 접근과 검토한 대안", "03 방법 선택")
tf = textbox(s, Cm(0.9), Cm(2.4), Cm(32), Cm(2.6))
bullets(tf, [
    (0, "채택: 실측 역산(초기값) + 후보 스윕(배치 렌더) + 픽셀 정량 비교 — '수동 좌표하강'", {"size": 15.5, "bold": True, "color": GREEN}),
    (1, "초기값: 원본 주사위 픽셀 크기(88px)→카메라-주사위 거리 25~30cm 역산, 웹캠 FOV 통념(60~70°)으로 범위 한정", {"size": 13}),
    (1, "1회 Isaac 부팅에 후보 3~4개 동시 렌더 → 주사위 화면 좌표(%)로 정량 채점 → 3라운드에 수렴", {"size": 13}),
])
rows = [
    ["대안", "내용", "기각 이유"],
    ["PnP / 체커보드 캘리브레이션",
     "대응점으로 카메라 외부·내부 파라미터를 해석적으로 풀기",
     "원본 씬에 캘리브레이션 타깃이 없고, 로봇 자체가 다름(OMX→M1013) — 정확한 해가 존재하지 않는 문제"],
    ["미분가능 렌더링 최적화",
     "렌더-원본 픽셀 손실을 경사하강으로 자동 최소화",
     "씬 내용이 달라(사람·로봇 외형) 픽셀 손실이 의미 없음 + 구축 비용 과다 — 7개 파라미터에 과한 도구"],
    ["Isaac GUI 수동 시점 조작",
     "창을 띄워 마우스로 맞추고 파라미터 읽기",
     "재현성 없음 · 정량 기록 안 남음 · FOV(내부 파라미터) 조정 불가에 가까움"],
    ["손목캠 롤 해석적 유도",
     "Camera 클래스 축 관례로 up 벡터를 수식으로 계산",
     ("시도했으나 실패 — up 인자와 실효 이미지-up 매핑이 관례와 달랐음 → 90° 스윕 실증으로 전환 (T/S ④)", {"color": ORANGE})],
]
table(s, Cm(0.9), Cm(5.6), Cm(32.1), rows, widths=[5.2, 8.2, 11.6], size=12, rh=Cm(1.7))
tf = textbox(s, Cm(0.9), Cm(15.0), Cm(32), Cm(2.8))
bullets(tf, [
    (0, "선택 기준: 완전 자동화가 불가능한 문제(정답 구도가 애초에 존재하지 않음)에서는, 사람이 판단 루프에 남되", {"size": 13.5}),
    (1, "렌더 배치·정량 채점·기록을 자동화해 반복 비용을 최소화하는 절충이 가장 빠르고 재현 가능했음", {"size": 13.5, "bold": True}),
])
foot(s)

# ═══════════ 7. 정면캠 3라운드 ═══════════
s = slide()
bar(s, "정면캠 탐색 — 3라운드 수렴 과정", "04 정면캠")
# 세 그리드의 렌더 높이를 5.2cm 로 통일 (R1/R2 는 2:1, R3 는 4:3 비율)
gh = Cm(5.2)
imgs = [("probe_grid.png", gh * 2, "R1 (A~D): 너무 멀다 — 상판 옆면·바닥 노출"),
        ("probe_grid2.png", gh * 2, "R2 (E~H): 테이블 앞으로 당김 — F 채택"),
        ("probe_grid3.png", gh * 4 / 3, "R3 (F1~F3): 흰벽+회색 테이블 — F1 확정")]
x = Cm(0.9)
for f, w_, cap_ in imgs:
    pic(s, os.path.join(IMG, f), x, Cm(2.5), w_)
    caption(s, x, Cm(2.5) + gh + Cm(0.08), w_, cap_, 11, True, INK)
    x = x + w_ + Cm(0.6)
tf = textbox(s, Cm(0.9), Cm(8.7), Cm(32), Cm(4.6))
bullets(tf, [
    (0, "R1→R2 판단: 원본은 카메라가 상판 바로 위(+10cm)에서 표면을 스치듯 봄 — 렌더는 멀어서 옆면·바닥까지 보임 → x 2.2→1.2~1.4m", {"size": 13}),
    (0, "R2→R3 판단: 구도는 맞았으나 배경(파란 격자 바닥)·테이블색(갈색)이 실물(흰 벽·회색)과 상이 → 씬 보정 (도메인 갭 직접 감소)", {"size": 13}),
    (0, "정량 채점 (주사위 화면 가로 위치): 원본 파랑 19% · 노랑 84%  →  F1 파랑 29% · 노랑 79% (후보 중 최근접)", {"size": 13.5, "bold": True}),
    (0, "확정 F1: pos (1.20, −0.25, 0.46) · look_at (0.45, −0.15, 0.52) · focal 15mm (HFOV 69.9°) · 640×480", {"size": 13.5, "bold": True, "color": GREEN, "mono": True}),
])
rows = [
    ["트러블슈팅 ③", "증상: 1차 프로브 전멸 (zero-norm quaternion 예외)"],
    ["원인", "look-at→쿼터니언 단순 공식(w=√(1+trace)/2 로 나눔)은 카메라가 정확히 −x 를 볼 때 180° 회전 → w=0 → 0-나눗셈"],
    ["해결", ("scipy Rotation.from_matrix 로 교체 (특이점 안전) — 기존 capture_m1013.py 는 카메라가 비스듬해 우연히 안 걸렸던 것", {"bold": True})],
]
table(s, Cm(0.9), Cm(13.7), Cm(32.1), rows, widths=[3.4, 18.6], size=12, header=False, rh=Cm(1.15))
foot(s)

# ═══════════ 8. 손목캠 — 거동 분석과 기하 역산 ═══════════
s = slide()
bar(s, "손목캠 — 실물 거동 분석 → 장착 기하 역산", "05 손목캠")
iw4 = Cm(14.0)
pic(s, os.path.join(IMG, "c1ep80_time.png"), Cm(0.9), Cm(2.5), iw4)
caption(s, Cm(0.9), Cm(2.5) + iw4 * 480 / 960 + Cm(0.08), iw4,
        "ep80 camera1 시간축 (80프레임 간격) — 주사위가 위에서 진입해 하단에서 커짐", 11.5, True, INK)
tf = textbox(s, Cm(15.6), Cm(2.5), Cm(17.4), Cm(8.2))
bullets(tf, [
    (0, "실물 프레임에서 읽어낸 3가지 (역산의 근거):", {"size": 14.5, "bold": True}),
    (1, "① 하향 시점 + 이미지 위 = 이동(접근) 방향 — 주사위가 위에서 진입하므로", {"size": 13}),
    (1, "② 파지점이 화면 하단 중앙 → 카메라가 파지축에서 접근 방향으로 ~6.5cm 오프셋 장착 (거리×tan 편차각 역산)", {"size": 13}),
    (1, "③ 파지 시 주사위 폭 ~250px → 카메라-주사위 6~10cm (OMX 는 캠이 핑거 근처)", {"size": 13}),
    (0, "sim 이식 (플랜지 로컬 좌표로 환산):", {"size": 14.5, "bold": True}),
    (1, "접근 방향 = 플랜지 로컬 (−0.655, −0.755, 0) — flange0 궤적에서 역산", {"size": 13, "mono": True}),
    (1, "전방 = 로컬 +z (손목 45° 보정 덕에 파지 시 세계 수직 하향)", {"size": 13}),
    (1, "부착 = link_6 의 자식 prim → USD 변환 상속으로 물리를 자동 추종 (조인트 불필요)", {"size": 13, "bold": True}),
])
tf = textbox(s, Cm(0.9), Cm(11.6), Cm(32), Cm(5.8))
bullets(tf, [
    (0, "M1013 스케일 제약: 그리퍼가 길어(플랜지→팁 0.21m) 실물만큼 주사위를 크게 볼 수 없음 → FOV 60°로 절충 (실물 대비 약 70% 크기)", {"size": 13.5}),
    (0, "라운드 구성: R1 (W1~4: 오프셋·FOV 스윕) → R2 (롤 90° 단위 실증) → R3 (V1~3: 롤 미세보정 +40°/55° · 오프셋 40/60mm)", {"size": 13.5}),
])
foot(s)

# ═══════════ 9. 트러블슈팅 ④ 롤 실증 ═══════════
s = slide()
bar(s, "트러블슈팅 ④ — 이미지 롤이 90° 틀어짐 → 실증 스윕으로 보정", "05 손목캠")
iw5 = Cm(13.2)
pic(s, os.path.join(IMG, "wprobe_grasp.png"), Cm(0.9), Cm(2.5), iw5)
caption(s, Cm(0.9), Cm(2.5) + iw5 * 480 / 960 + Cm(0.08), iw5,
        "R1: 큐브가 하단이 아닌 오른쪽에 — up 축 매핑이 가정과 다름", 11.5, True, INK)
pic(s, os.path.join(IMG, "wprobe2_grasp.png"), Cm(0.9), Cm(10.3), iw5)
caption(s, Cm(0.9), Cm(10.3) + iw5 * 480 / 960 + Cm(0.08), iw5,
        "R2: up 인자 90° 단위 스윕 — 화면이 함께 회전함을 실증 (L90 우상 → 180 좌상 → R90 좌하)", 11.5, True, INK)
tf = textbox(s, Cm(14.8), Cm(2.5), Cm(18.2), Cm(14.6))
bullets(tf, [
    (0, "증상: 오프셋을 '이미지 위' 방향으로 줬는데 큐브가 순수하게 '오른쪽'으로 밀림", {"size": 14}),
    (0, "진단 논리 (기하로 모순 증명):", {"size": 14, "bold": True}),
    (1, "오프셋 o 를 up 방향으로 주면 파지점의 화면 변위는 up 성분만 가질 수 있음 (right 성분 = o·(f×up) = 0)", {"size": 12.5}),
    (1, "그런데 순수 right 변위가 관측됨 → 실효 이미지-up 이 지정한 up 과 90° 어긋났다는 뜻", {"size": 12.5, "bold": True}),
    (0, "원인(추정): Isaac Camera 의 world-axes 규약에서 up 인자→실효 롤 매핑이 look-at 축 조합에 따라 문서 통념과 다름", {"size": 13}),
    (0, "해결: 해석 포기 → 실증 캘리브레이션", {"size": 14, "bold": True, "color": GREEN}),
    (1, "R2: up 인자를 90° 단위로 4방향 렌더 → 'up 인자 +90° = 화면 +90° 회전' 관계와 현재 오차각(215° vs 목표 270°)을 실측", {"size": 12.5}),
    (1, "R3: 부족분 +40°/+55° 후보로 미세보정 → +55°(R145)가 하단 중앙 적중", {"size": 12.5, "bold": True}),
    (0, "기록: 확정 up_arg 는 '보정된 값'이므로 JSON 에 경고 주석과 함께 저장 — 재유도 금지, 그대로 쓸 것", {"size": 13, "color": ORANGE}),
])
foot(s)

# ═══════════ 10. 손목캠 확정 ═══════════
s = slide()
bar(s, "손목캠 확정 (V3) + 트러블슈팅 ⑤ 테이블 확대", "05 손목캠")
iw6 = Cm(10.8)
pic(s, os.path.join(IMG, "wprobe3_grasp.png"), Cm(0.9), Cm(2.5), iw6)
caption(s, Cm(0.9), Cm(2.5) + iw6 * 3 / 4 + Cm(0.08), iw6,
        "R3 파지 컷 — V3: 큐브가 실물처럼 하단(38%, 86%)에서 프레임에 걸림", 11.5, True, INK)
pic(s, os.path.join(IMG, "wprobe3_home.png"), Cm(12.2), Cm(2.5), iw6)
caption(s, Cm(12.2), Cm(2.5) + iw6 * 3 / 4 + Cm(0.08), iw6,
        "R3 홈 컷 — 회색 표면 + 핑거 가장자리 (실물 f0 과 같은 구성)", 11.5, True, INK)
tf = textbox(s, Cm(23.4), Cm(2.5), Cm(9.6), Cm(9))
bullets(tf, [
    (0, "확정 V3:", {"size": 14, "bold": True, "color": GREEN}),
    (1, "local_pos (−0.039, −0.045, 0.05)", {"size": 12, "mono": True}),
    (1, "forward = 로컬 +z", {"size": 12, "mono": True}),
    (1, "up_arg (−0.970, −0.243, 0)", {"size": 12, "mono": True}),
    (1, "focal 18mm (HFOV 60.3°)", {"size": 12, "mono": True}),
    (0, "정량: 파지 컷 큐브 중심 (0.38, 0.86) · 폭 170px — 실물 (~0.33, ~0.9) · ~250px", {"size": 12.5}),
])
rows = [
    ["트러블슈팅 ⑤", "증상: 파지 컷에 실물에 없는 테이블 모서리·파란 바닥이 보임"],
    ["원인", "M1013 파지점이 테이블 구석에서 13~14cm — 하향 60~70° 시야에 모서리 진입 (실물은 시야 대비 테이블이 큼)"],
    ["해결", ("렌더용 테이블만 (0.49,0.60)→(0.85,1.10)m 확대, 상판 z 동일 — 물리 검증 씬은 원치수 유지. 결과: 모서리 완전 소거", {"bold": True})],
]
table(s, Cm(0.9), Cm(12.4), Cm(32.1), rows, widths=[3.4, 18.6], size=12, header=False, rh=Cm(1.15))
tf = textbox(s, Cm(0.9), Cm(16.3), Cm(32), Cm(1.6))
bullets(tf, [
    (0, "확대는 물리에 무해: 큐브 pick/place 는 테이블 중앙부라 지지면 변화 없음 — 단, 5-3 씬 조립 시 이 값(scene_render)을 쓸 것", {"size": 12.5, "color": ORANGE}),
])
foot(s)

# ═══════════ 11. 최종 비교 시트 ═══════════
s = slide()
bar(s, "최종 결과 — 실물 vs sim 4쌍 비교 (cam_match_v1.png)", "06 결과")
mh = Cm(15.6)
mw = mh * 670 / 1000
if os.path.exists(os.path.join(SIMOUT, "cam_match_v1.png")):
    s.shapes.add_picture(os.path.join(SIMOUT, "cam_match_v1.png"), Cm(0.9), Cm(2.3), height=mh)
tf = textbox(s, Cm(0.9) + mw + Cm(0.8), Cm(2.7), SW - mw - Cm(2.6), Cm(15))
bullets(tf, [
    (0, "왼쪽 열 = OMX 실물 학습 영상, 오른쪽 열 = M1013 Isaac 렌더 (같은 시점: 홈 / 파지)", {"size": 14.5, "bold": True}),
    (0, "정합된 것 (구도):", {"size": 14, "bold": True, "color": GREEN}),
    (1, "정면캠: 낮은 앵글 · 테이블이 하단 채움 · pick 좌 / place+노랑 우 · 그리퍼 중앙", {"size": 13}),
    (1, "손목캠: 하향 + 이미지 위=접근 방향 · 파지점 하단 중앙 · 핑거 가장자리", {"size": 13}),
    (1, "배경: 흰 벽 · 회색 테이블 (실물 근사)", {"size": 13}),
    (0, "남는 차이 = 5-3 에서 측정할 순수 도메인 갭:", {"size": 14, "bold": True, "color": ORANGE}),
    (1, "로봇 외형 (검정 OMX ↔ 흰색 M1013 · 크기)", {"size": 13}),
    (1, "사람·의자·파티션 등 배경 사물 없음", {"size": 13}),
    (1, "질감·조명 (실사 노이즈 ↔ 클린 렌더)", {"size": 13}),
    (1, "주사위 외형 (둥근 폼 주사위 ↔ 각진 큐브)", {"size": 13}),
    (0, "이 차이들이 closed-loop 성능에 주는 영향이 크면 → 5-4 co-training(159ep 재렌더 혼합)으로 대응", {"size": 13.5}),
])
foot(s)

# ═══════════ 12. 생성 파일 카탈로그 ═══════════
s = slide()
bar(s, "오늘 생성한 파일 전체 — 무엇이 무엇인지", "07 산출물")
rows = [
    ["파일 (BASE=/home/kim/m1013)", "종류", "내용 · 용도"],
    [("cameras_v6match.json", {"mono": True, "bold": True}), "확정값",
     ("★ 카메라 2대 확정 파라미터 + 렌더 씬 설정(확대 테이블·벽) — 5-3 렌더러가 읽을 단일 소스", {"bold": True})],
    [("frontcam_probe.py", {"mono": True}), "코드",
     "정면캠 후보 배치 렌더 (CANDS 리스트 수정→재실행). 확정 F1 후보가 코드에 남아있음"],
    [("wristcam_probe.py", {"mono": True}), "코드",
     "손목캠 후보 배치 렌더 — link_6 부착·롤 실증·홈/파지 2포즈. 확정 V3 = V3_r145_off60"],
    [("sim_out/cam_match_v1.png", {"mono": True, "bold": True}), "이미지",
     ("★ 실물 vs sim 4쌍 비교시트 — 보고·랩미팅용 최종 결과물", {"bold": True})],
    [("sim_out/frontcam_probe/*.png", {"mono": True}), "이미지", "정면캠 후보별 렌더 (R1 A~D · R2 E~H · R3 F1~F3)"],
    [("sim_out/wristcam_probe/*.png", {"mono": True}), "이미지", "손목캠 후보별 렌더 (이름_home / 이름_grasp 2포즈씩)"],
    [("대화록 및 PPT/img_0824/*.png", {"mono": True}), "이미지",
     "분석 근거 자료 — 원본 프레임(ref_*)·거동 그리드(wrist_ref, c1ep80)·라운드별 비교(probe_grid*, wprobe*)"],
    [("make_report_pptx_0824.py", {"mono": True}), "코드", "이 PPT 생성 스크립트 (남색 바+flowbox 체계)"],
]
table(s, Cm(0.9), Cm(2.6), Cm(32.1), rows, widths=[7.2, 2.2, 12.6], size=12, rh=Cm(1.35))
tf = textbox(s, Cm(0.9), Cm(15.6), Cm(32), Cm(2.6))
bullets(tf, [
    (0, "재실행 (프로브): cd /home/kim/isaacsim && ./python.sh /home/kim/m1013/frontcam_probe.py  (wristcam 동일 — 회당 ~1분)", {"size": 12.5, "mono": True}),
    (0, "실물 기준 프레임 재추출: docker exec physical_ai_server ffmpeg … (경로는 T/S ①② 슬라이드 참조)", {"size": 12.5, "mono": True}),
])
foot(s)

# ═══════════ 13. 다음 시작점 — 5-3 closed-loop ═══════════
s = slide()
bar(s, "다음 시작점 — STEP 5-3 closed-loop 첫 가동", "08 다음")
loop2 = [
    ("① 렌더 루프 조립\nreplay_isaac 씬\n+ JSON 카메라 2대", NAVY),
    ("② 관측 파이프\n매 스텝 640×480 2뷰\n+ state(관절 실측)", NAVY),
    ("③ 정책 추론 연결\n컨테이너 ACT 체크포인트\n↔ 호스트 Isaac 브릿지", ORANGE),
    ("④ 실행·평가\n성공률 · 오차\nvs open-loop 대비", NAVY),
    ("⑤ 갭 정량화\n실패 모드 분석\n→ 5-4 여부 결정", GREEN),
]
bw, bh, gap = Cm(5.85), Cm(2.5), Cm(0.7)
x0 = Cm(0.9)
for i, (txt, col) in enumerate(loop2):
    x = x0 + i * (bw + gap)
    flowbox(s, x, Cm(2.6), bw, bh, txt, fill=col, size=12)
    if i < len(loop2) - 1:
        arrow(s, x + bw + Cm(0.0), Cm(2.6) + bh / 2 - Cm(0.25))
tf = textbox(s, Cm(0.9), Cm(5.7), Cm(32), Cm(7))
bullets(tf, [
    (0, "구체 시작점: closed_loop_m1013.py 신규 작성 — replay_isaac.py 의 씬 조립부 재사용 + cameras_v6match.json 로드", {"size": 14, "bold": True}),
    (1, "관측 규약은 학습과 동일하게: observation.images.camera1=손목캠, camera2=정면캠, state=action[t−1] 규약 확인 필요", {"size": 13}),
    (0, "고려사항 (착수 전 결정할 것):", {"size": 14, "bold": True}),
    (1, "㉠ 추론 위치: 정책은 컨테이너(lerobot·GPU), Isaac 은 호스트 — 파일 핑퐁(느림) vs ZMQ/소켓 브릿지(권장) vs 호스트에 lerobot 설치", {"size": 13}),
    (1, "㉡ 실행 방식: 30Hz 매 스텝 추론 불필요 — ACT 청크(예: 100프레임) 단위 실행 + 재계획 주기 결정", {"size": 13}),
    (1, "㉢ 렌더 속도: 두 카메라 렌더 포함 스텝 속도 실측 → 실시간 불가 시 sim 시간 기준으로 평가 (물리 정합성 유지)", {"size": 13}),
    (1, "㉣ 평가 설계: 동일 에피소드 GT 재생(6/8 성공)과 같은 8ep 로 A/B — '구도 정합 완료' 덕에 실패 원인을 도메인 갭으로 귀속 가능", {"size": 13}),
])
rows = [
    ["병행·후속 (0820 계획 유지)", "내용"],
    ["5-4 co-training A/B", "5-3 갭이 크면: 159ep 를 확정 카메라로 재렌더 → sim 혼합학습 이득 검증 후 채택"],
    ["실기 준비", "TCP 0.12m 잠정값 (LEHR 실측 후 재변환 1분+재학습 ~3h) · 어태치먼트 연질 팁 · 손목캠 마운트 설계에 오늘 오프셋(−39,−45,+50mm) 참조 가능"],
]
table(s, Cm(0.9), Cm(13.6), Cm(32.1), rows, widths=[5.2, 16.8], size=12, rh=Cm(1.2))
foot(s)

# ═══════════ 14. 링크·참조 ═══════════
s = slide()
bar(s, "링크 · 참조 경로", "09 참조")
rows = [
    ["자원", "위치 / URL"],
    ["학습 데이터셋 (M1013 v6)", ("https://huggingface.co/datasets/dlcodnjs/m1013_act_pick_and_place_v6_joint", {"mono": True})],
    ["원본 OMX 영상 (실체)", ("physical_ai_server 컨테이너: /root/.cache/huggingface/lerobot/dlcodnjs/omx_act_pick_and_place_v4_162/videos/", {"mono": True})],
    ["M1013 코드 저장소", ("https://github.com/dlcodnjs817/M1013  (비공개 — 오늘 파일 커밋 예정)", {"mono": True})],
    ["TX90 저장소 (선행 작업)", ("https://github.com/dlcodnjs817/tx90  (비공개)", {"mono": True})],
    ["M1013 USD·URDF 출처", ("https://github.com/DoosanRobotics/doosan-robot2  (dsr_description2/usd/m1013.usd)", {"mono": True})],
    ["씬 좌표 단일 소스", ("/home/kim/m1013/replay_ep000.npz  (cube_pick/place · table_top_z · flange0 · t_close=226)", {"mono": True})],
    ["카메라 확정값", ("/home/kim/m1013/cameras_v6match.json", {"mono": True, "bold": True})],
    ["지난 보고 (선행 맥락)", ("tx90/대화록 및 PPT/업무보고_20260820.md · M1013_20260820.pptx · 랩미팅 make_labmeeting_pptx_0824.py", {"mono": True})],
]
table(s, Cm(0.9), Cm(2.7), Cm(32.1), rows, widths=[5.4, 16.6], size=12, rh=Cm(1.25))
tf = textbox(s, Cm(0.9), Cm(14.2), Cm(32), Cm(3.6))
bullets(tf, [
    (0, "한 줄 요약: 정책이 학습 때 본 두 화면(손목캠·정면캠)을 Isaac 에 복원 완료 — 파라미터는 JSON 에 고정,", {"size": 14.5, "bold": True}),
    (1, "다음 세션은 closed_loop_m1013.py 작성부터 (관측 규약 확인 → 추론 브릿지 결정 → 8ep A/B)", {"size": 13.5}),
])
foot(s)

out = os.path.join(OUTDIR, "M1013_카메라정합_20260824.pptx")
prs.save(out)
print("저장:", out)
